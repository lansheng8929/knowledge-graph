"""parsers 单元测试（无外部依赖）。"""

from io import BytesIO

import pytest
from openpyxl import Workbook

from app.models import ParseError
from app.parsers import get_parser


def _xlsx(entities_rows, edges_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "entities"
    ws.append(["id", "nodeType", "label", "gender"])
    for r in entities_rows:
        ws.append(r)
    ws2 = wb.create_sheet("edges")
    ws2.append(["source", "target", "linkType", "amount"])
    for r in edges_rows:
        ws2.append(r)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_csv_entities():
    raw = b"id,nodeType,label,gender\np1,person,\xe5\xbc\xa0\xe4\xb8\x89,\xe7\x94\xb7\n"
    parser = get_parser("people.csv")
    tables = parser.parse(raw, "people.csv", kind="entities")
    assert list(tables.keys()) == ["entities"]
    assert tables["entities"][0]["id"] == "p1"
    assert tables["entities"][0]["gender"] == "\u7537"


def test_csv_edges_kind():
    raw = b"source,target,linkType,amount\np1,p2,TRANSFER,5000\n"
    parser = get_parser("x.csv")
    tables = parser.parse(raw, "x.csv", kind="edges")
    assert list(tables.keys()) == ["edges"]
    assert tables["edges"][0]["linkType"] == "TRANSFER"


def test_csv_skips_blank_rows():
    raw = b"a,b\n1,2\n\n3,4\n"
    tables = get_parser("x.csv").parse(raw, "x.csv")
    assert len(tables["entities"]) == 2


def test_excel_both_sheets():
    raw = _xlsx(
        [["p1", "person", "张三", "男"]],
        [["p1", "p2", "TRANSFER", 5000]],  # 数字应保留类型
    )
    tables = get_parser("book.xlsx").parse(raw, "book.xlsx")
    assert set(tables.keys()) == {"entities", "edges"}
    assert tables["entities"][0]["nodeType"] == "person"
    assert tables["edges"][0]["amount"] == 5000


def test_excel_skips_empty_rows_and_dupe_headers():
    raw = _xlsx([["p1", "person", "张三", "男"], [None, None, None, None]], [])
    # 仅 entities sheet 有数据
    tables = get_parser("book.xlsx").parse(raw, "book.xlsx")
    assert len(tables["entities"]) == 1


def test_excel_sheet_mapping():
    """用户指定 sheet 名（非 entities/edges）也能抽取。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append(["id", "nodeType", "label"])
    ws.append(["p1", "person", "张三"])
    ws2 = wb.create_sheet("转账关系")
    ws2.append(["source", "target", "linkType"])
    ws2.append(["p1", "p2", "OWNS"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    raw = buf.getvalue()

    tables = get_parser("b.xlsx").parse(
        raw, "b.xlsx", sheet_mapping={"entities": "人员", "edges": "转账关系"}
    )
    assert set(tables.keys()) == {"entities", "edges"}
    assert tables["entities"][0]["nodeType"] == "person"
    assert tables["edges"][0]["linkType"] == "OWNS"


def test_excel_positional_default():
    """默认按顺序：第一个非空表=entities，第二个=edges（不依赖 sheet 名）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append(["id", "nodeType", "label"])
    ws.append(["p1", "person", "张三"])
    ws2 = wb.create_sheet("转账关系")
    ws2.append(["source", "target", "linkType"])
    ws2.append(["p1", "p2", "OWNS"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    tables = get_parser("b.xlsx").parse(raw=buf.getvalue(), filename="b.xlsx")
    assert list(tables.keys()) == ["entities", "edges"]
    assert tables["entities"][0]["nodeType"] == "person"
    assert tables["edges"][0]["linkType"] == "OWNS"


def test_excel_positional_skips_empty_sheets():
    """空表不占位：前面有空 sheet 时，第一个有数据的仍是 entities。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"  # 空表
    ws2 = wb.create_sheet("数据1")
    ws2.append(["id", "nodeType", "label"])
    ws2.append(["p1", "person", "张三"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    tables = get_parser("b.xlsx").parse(raw=buf.getvalue(), filename="b.xlsx")
    assert list(tables.keys()) == ["entities"]
    assert tables["entities"][0]["nodeType"] == "person"


def test_unsupported_extension():
    with pytest.raises(ParseError):
        get_parser("x.docx")

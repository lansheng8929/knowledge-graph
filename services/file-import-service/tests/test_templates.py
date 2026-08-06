"""模板引擎单元测试：列重映射、选表、CJK 标识符、模板端点。"""

import io
import json

from fastapi.testclient import TestClient

from app.main import app
from app.templates.builtin import BUILTIN_TEMPLATES
from app.templates.engine import get_template, parse_with_template

client = TestClient(app)

ADMIN = json.dumps(
    {
        "username": "admin",
        "uid": "u-admin",
        "tenantId": "default",
        "clearance": 3,
        "roles": ["admin", "analyst", "privileged"],
        "teams": ["ops"],
    },
    ensure_ascii=False,
)


def _csv_bytes(text: str) -> bytes:
    return text.encode("utf-8")


def _config_json(**over):
    cfg = {
        "tags": {
            "tenantId": "default",
            "classification": 0,
            "owner": "admin",
            "visibility": "internal",
        },
        "edge": {"insertMode": "update"},
    }
    cfg.update(over)
    return json.dumps(cfg, ensure_ascii=False)


# ── 引擎：通话关系模板（单 CSV，边规则自动建 phone 实体）────────


def test_call_template_remaps_columns():
    call = _csv_bytes(
        "拨通方,接听方,通话时间\n"
        "13800001111,13800002222,2026-08-01 10:00:00\n"
        "13800003333,13800001111,2026-08-01 11:00:00\n"
    )
    tpl = get_template("tpl-call")
    tables = parse_with_template(("calls.csv", call), None, tpl)
    # 模板同时配了实体+边抽取：实体从 拨通方/接听方 列抽（按值去重），边从同表映射
    assert "edges" in tables and "entities" in tables
    ids = [e["id"] for e in tables["entities"]]
    assert sorted(ids) == ["13800001111", "13800002222", "13800003333"]
    assert tables["entities"][0]["nodeType"] == "phone"
    row0 = tables["edges"][0]
    assert row0["source"] == "13800001111"
    assert row0["target"] == "13800002222"
    assert row0["linkType"] == "通话"
    # time 不由模板配置（固定为边的更新时间，程序写入），行内不再有 time
    assert "time" not in row0
    # 已映射的结构列（拨通方/接听方）被剔除，不混进 props
    assert "拨通方" not in row0 and "接听方" not in row0
    # 业务键列保留（props 只存 通话时间 等未消费列）
    assert row0["通话时间"]


def test_call_template_contains_column_match():
    # 列名带前缀，关键词包含匹配应命中
    call = _csv_bytes("拨通方号码,接听方号码,通话起始时间\nA,B,T1\n")
    tpl = get_template("tpl-call")
    tables = parse_with_template(("calls.csv", call), None, tpl)
    assert tables["edges"][0]["source"] == "A"
    assert tables["edges"][0]["target"] == "B"
    # 实体也从含前缀的列抽（A/B 各一个 phone）
    ids = sorted(e["id"] for e in tables["entities"])
    assert ids == ["A", "B"]


def test_trade_template_two_sheets_positional():
    # 单文件：第一个非空表=账户（entities），第二个=转账（edges）
    wb = io.BytesIO()
    from openpyxl import Workbook

    w = Workbook()
    ws1 = w.active
    ws1.title = "账户"
    ws1.append(["账号", "户名"])
    ws1.append(["acc-1", "张三"])
    ws1.append(["acc-2", "李四"])
    ws2 = w.create_sheet("转账记录")
    ws2.append(["转出账号", "转入账号", "转账时间"])
    ws2.append(["acc-1", "acc-2", "2026-08-01T09:00:00Z"])
    w.save(wb)

    tpl = get_template("tpl-trade")
    tables = parse_with_template(("trades.xlsx", wb.getvalue()), None, tpl)
    assert tables["entities"][0]["id"] == "acc-1"
    assert tables["entities"][0]["nodeType"] == "account"
    assert tables["edges"][0]["source"] == "acc-1"
    assert tables["edges"][0]["target"] == "acc-2"
    assert tables["edges"][0]["linkType"] == "转账"


def test_bare_template_has_no_rules():
    assert get_template("tpl-bare") is not None
    assert not (
        get_template("tpl-bare").get("entityRule")
        or get_template("tpl-bare").get("edgeRule")
    )


def test_unknown_template_none():
    assert get_template("nope") is None


# ── 端点 ─────────────────────────────────────────────


def test_templates_endpoint():
    r = client.get("/api/v1/import/templates", headers={"X-User-Context": ADMIN})
    assert r.status_code == 200
    data = r.json()["data"]
    assert data[0]["id"] == "tpl-call"
    assert data[0]["default"] is True
    assert any(t["id"] == "tpl-bare" for t in data)


def test_preview_with_call_template():
    call = _csv_bytes(
        "拨通方,接听方,通话时间\n13800001111,13800002222,2026-08-01 10:00:00\n"
    )
    r = client.post(
        "/api/v1/import/preview",
        headers={"X-User-Context": ADMIN},
        data={"config": _config_json(), "template_id": "tpl-call"},
        files={"file": ("calls.csv", call, "text/csv")},
    )
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    # 通话模板同时抽取实体+边：1 行通话 → 2 个 phone 实体（拨通方/接听方）+ 1 条「通话」边
    assert data["edgeCount"] == 1
    assert data["entityCount"] == 2
    assert data["edges"][0]["linkType"] == "通话"
    assert data["entities"][0]["nodeType"] == "phone"


def test_preview_bare_template_standard_headers():
    # 通用模板：标准表头，单文件 Excel 两张表（entities + edges）
    import io
    from openpyxl import Workbook

    wb = io.BytesIO()
    w = Workbook()
    ws1 = w.active
    ws1.title = "entities"
    ws1.append(["id", "nodeType", "label"])
    ws1.append(["p1", "person", "张三"])
    ws2 = w.create_sheet("edges")
    ws2.append(["source", "target", "linkType"])
    ws2.append(["p1", "p1", "OWNS"])
    w.save(wb)

    r = client.post(
        "/api/v1/import/preview",
        headers={"X-User-Context": ADMIN},
        data={"config": _config_json(), "template_id": "tpl-bare"},
        files={
            "file": (
                "data.xlsx",
                wb.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["data"]["entityCount"] == 1
    assert r.json()["data"]["edgeCount"] == 1


def test_bare_csv_edges_auto_detect():
    # 无模板：CSV 表头含 source/target/linkType → 自动判为边表
    edg = _csv_bytes("source,target,linkType\np1,p1,OWNS\n")
    ent = _csv_bytes("id,nodeType,label\np1,person,张三\n")
    r = client.post(
        "/api/v1/import/preview",
        headers={"X-User-Context": ADMIN},
        data={
            "config": _config_json(dangling="auto-create", danglingNodeType="person")
        },
        files={"file": ("r.csv", edg, "text/csv")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["data"]["edgeCount"] == 1
    r2 = client.post(
        "/api/v1/import/preview",
        headers={"X-User-Context": ADMIN},
        data={"config": _config_json()},
        files={"file": ("e.csv", ent, "text/csv")},
    )
    assert r2.status_code == 200, r2.text
    assert r2.json()["data"]["entityCount"] == 1


def test_cjk_identifier_allowed_by_validator():
    from app.validator import validate
    from app.ir import ParsedEdge, ParsedEntity, ParsedGraph
    from app.models import ImportConfig

    graph = ParsedGraph(
        entities=[ParsedEntity(id="p1", nodeType="person", label="张三")],
        edges=[ParsedEdge(id="e1", source="p1", target="p1", linkType="通话")],
    )
    res = validate(graph, ImportConfig())
    assert not res.errors


# ── 历史任务列表 ─────────────────────────────────────


def test_task_store_list_ordering():
    from app.tasks import TaskStore

    s = TaskStore()
    a = s.create("a.csv")
    b = s.create("b.csv")
    lst = s.list()
    assert lst[0].id == b.id  # 新的在前
    assert [t.filename for t in lst] == ["b.csv", "a.csv"]
    # 摘要不含 errors/warnings
    assert "errors" not in lst[0].summary_dict()


def test_tasks_list_endpoint():
    r = client.get("/api/v1/import/tasks", headers={"X-User-Context": ADMIN})
    assert r.status_code == 200
    data = r.json()["data"]
    assert isinstance(data, list)
    for item in data:
        assert "errors" not in item  # 列表为轻量摘要
        assert {"id", "status", "filename"}.issubset(item)


# ── 选表不做位置兜底（2026-08-05）─────────────────────


def test_mismatched_template_errors_no_positional_fallback():
    # 实体文件（标准表头）用「通话」模板 → 找不到 拨通/接听 表 → 报 400，
    # 不再把标准实体行位置兜底成通话边/实体。
    ent = _csv_bytes("id,nodeType,label\np1,person,张三\n")
    r = client.post(
        "/api/v1/import/preview",
        headers={"X-User-Context": ADMIN},
        data={"config": _config_json(), "template_id": "tpl-call"},
        files={"file": ("e.csv", ent, "text/csv")},
    )
    assert r.status_code == 400, r.text
    assert "模板未匹配" in r.json()["detail"]


def test_template_partial_match_surfaces_warning():
    # 边表匹配、实体表未匹配 → 警告带模板提示，且不做位置兜底
    # （不会把「转出/转入」行硬当成 account 实体）
    data = _csv_bytes("转出,转入,时间\nacc-1,acc-2,T1\n")
    r = client.post(
        "/api/v1/import/preview",
        headers={"X-User-Context": ADMIN},
        data={
            "config": _config_json(dangling="auto-create", danglingNodeType="account"),
            "template_id": "tpl-trade",
        },
        files={"file": ("t.csv", data, "text/csv")},
    )
    assert r.status_code == 200, r.text
    body = r.json()["data"]
    # 实体规则未匹配 → 不位置兜底建实体
    assert body["entityCount"] == 0
    assert any("未匹配到实体表" in w for w in body["warnings"])

"""Excel(.xlsx/.xlsm) 解析：`entities` / `edges` 两个 sheet（完美格式化假设）。"""

from __future__ import annotations

import io
from typing import Dict, List

from openpyxl import load_workbook

from . import Parser, Table, norm_value, register

RESERVED_SHEETS = ("entities", "edges")


def _sheet_rows(ws) -> List[Dict[str, object]]:
    """首行=表头，其余=数据行；跳过空行；重复表头后缀去重。"""
    headers: List[str] = []
    rows: List[Dict[str, object]] = []
    seen: Dict[str, int] = {}
    for record in ws.iter_rows(values_only=True):
        if not record or all(c is None or str(c).strip() == "" for c in record):
            continue
        if not headers:
            for c in record:
                h = str(c or "").strip()
                if h:
                    seen[h] = seen.get(h, 0) + 1
                    h = h if seen[h] == 1 else f"{h}_{seen[h]}"
                headers.append(h)
            continue
        row: Dict[str, object] = {}
        for idx, h in enumerate(headers):
            if not h or idx >= len(record):
                continue
            row[h] = norm_value(record[idx])
        if any(str(v).strip() != "" for v in row.values()):
            rows.append(row)
    return rows


class ExcelParser(Parser):
    name = "excel"
    extensions = (".xlsx", ".xlsm")

    def parse(
        self,
        raw: bytes,
        filename: str,
        kind: str = "",
        sheet_mapping: dict = None,
    ) -> Table:
        """读取 entities/edges 两个 sheet。

        默认行为：按 sheet 顺序匹配——第一个非空表 = entities，第二个 = edges；
        用户可经 `sheet_mapping`（{ "entities": "人员表", "edges": "转账关系" }）指定。
        """
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            # 先抽所有非空表（保持 sheet 顺序；空表不占位）
            non_empty: Dict[str, List[Dict[str, object]]] = {}
            for sheet_name in wb.sheetnames:
                rows = _sheet_rows(wb[sheet_name])
                if rows:
                    non_empty[sheet_name] = rows

            tables: Table = {}
            if sheet_mapping:
                # 用户指定：按配置精确取表
                for target in RESERVED_SHEETS:
                    name = sheet_mapping.get(target)
                    if name and name in non_empty:
                        tables[target] = non_empty[name]
            else:
                # 默认：第一个非空表=entities，第二个=edges
                for idx, (_, rows) in enumerate(non_empty.items()):
                    if idx >= len(RESERVED_SHEETS):
                        break
                    tables[RESERVED_SHEETS[idx]] = rows
            return tables
        finally:
            wb.close()

    def parse_sheets(
        self, raw: bytes, filename: str
    ) -> Dict[str, List[Dict[str, object]]]:
        """按 sheet 名返回所有非空表（供模板引擎自主选表）。"""
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            non_empty: Dict[str, List[Dict[str, object]]] = {}
            for sheet_name in wb.sheetnames:
                rows = _sheet_rows(wb[sheet_name])
                if rows:
                    non_empty[sheet_name] = rows
            return non_empty
        finally:
            wb.close()


register(ExcelParser())  # noqa: E402
